from pptx import Presentation
import sys

try:
    from pptx import Presentation
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx", "-q"])
    from pptx import Presentation

prs = Presentation(r'c:\Users\iampr\Fathom\FATHOM.pptx')
for i, slide in enumerate(prs.slides, 1):
    print(f'=== SLIDE {i} ===')
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    print(text)
    print()
